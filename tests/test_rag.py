"""RAG contract tests. Embedding/rerank are fixtures; Qdrant here is embedded."""
from __future__ import annotations

import hashlib
import io
from unittest.mock import AsyncMock

import pytest
import pytest_asyncio

from opencmo import storage
from opencmo.rag import ingestion, providers, store
from opencmo.rag.config import RagSettings
from opencmo.rag.integration import content_request, safe_handoff_filter
from opencmo.rag.parsing import ParseError, parse_bytes, structured_text
from opencmo.rag.retrieval import fuse, retrieve, validate_output
from opencmo.rag.splitting import split_document
from opencmo.rag.text import tokens
from opencmo.rag.types import RetrievalRequest


@pytest_asyncio.fixture
async def library(tmp_path, monkeypatch):
    from qdrant_client import AsyncQdrantClient
    monkeypatch.setattr(storage, "_DB_PATH", tmp_path / "rag.db")
    monkeypatch.setenv("OPENCMO_RAG_STORAGE_PATH", str(tmp_path / "files"))
    monkeypatch.setenv("OPENCMO_RAG_ENABLED", "1")
    client = AsyncQdrantClient(":memory:")
    embed_impl = providers.embed
    close = client.close
    monkeypatch.setattr(client, "close", AsyncMock())
    monkeypatch.setattr(providers.vectors, "client", lambda: client)
    async def embedding(texts, settings, **kwargs):
        result = []
        for text in texts:
            vector = [0.01] * 8
            for term in text.lower().split():
                vector[int(hashlib.sha256(term.encode()).hexdigest(), 16) % 8] += 1
            result.append(vector)
        return result
    async def ranking(query, documents, settings, **kwargs):
        return sorted([(i, .95 if any(term in text.lower() for term in query.lower().split()) else .3)
                       for i, text in enumerate(documents)], key=lambda x: -x[1])
    monkeypatch.setattr(providers, "embed", embedding)
    monkeypatch.setattr(providers, "rerank", ranking)
    from opencmo.rag import retrieval
    monkeypatch.setattr(retrieval, "plan_queries", AsyncMock(side_effect=lambda req: [req.query]))
    user, account = await storage.create_user_with_account("rag-a@example.test", "password123")
    other_user, other = await storage.create_user_with_account("rag-b@example.test", "password123")
    a, b = account["id"], other["id"]
    p = await storage.ensure_project("Product", "https://example.com", "notes", account_id=a)
    p2 = await storage.ensure_project("Other", "https://other.example.com", "notes", account_id=a)
    foreign = await storage.ensure_project("Foreign", "https://foreign.example.com", "notes", account_id=b)
    settings = RagSettings(enabled=True, embedding_api_key="test-embedding", rerank_api_key="test-rerank")
    for account_id in (a, b):
        await storage.set_account_setting(account_id, "RAG_CONFIG", settings.model_dump_json())
    yield {"account": a, "other": b, "user": user['id'], "other_user": other_user['id'], "project": p, "project2": p2, "foreign": foreign, "settings": settings, 'embed_impl': embed_impl}
    await close()


async def add(lib, text, **kwargs):
    result = await ingestion.create_document(lib["account"], project_id=lib["project"],
        title=kwargs.pop("title", "Customer evidence"), data=text.encode(), **kwargs)
    if not result.get("duplicate"):
        await ingestion.ingest(lib["account"], result["document_id"], result["version_id"])
    return result


@pytest.mark.parametrize("text", [
    "# 标题\n\n客户星河公司的转化率为 12%。\n\n## 结果\n" + "增长效果明显，成本降低。" * 400,
    "# English\n\n" + "Customers reduced acquisition costs. " * 600,
    "# Table\n\n| Client | Revenue |\n| --- | --- |\n" + "| Acme | 42 |\n" * 200,
    "# Code\n\n" + chr(96) * 3 + "python\n" + "print('example')\n" * 150 + chr(96) * 3,
])
def test_parent_child_source_spans_and_budgets(text):
    settings = RagSettings()
    parsed = structured_text(text)
    chunks = split_document(parsed, "stable-generation", settings)
    parents = {c.id: c for c in chunks if c.level == "parent"}
    assert chunks
    assert [c.id for c in chunks] == [c.id for c in split_document(parsed, "stable-generation", settings)]
    for chunk in chunks:
        assert chunk.text == parsed.text[chunk.start:chunk.end]
        assert tokens(chunk.text) <= (settings.parent_tokens if chunk.level == "parent" else settings.child_tokens)
        if chunk.parent_id:
            parent = parents[chunk.parent_id]
            assert parent.start <= chunk.start < chunk.end <= parent.end


def test_docx_parser_keeps_headings_and_tables():
    from docx import Document
    document = Document()
    document.add_heading("客户案例", level=1)
    document.add_paragraph("星河客户使用产品。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "客户", "收益"
    table.cell(1, 0).text, table.cell(1, 1).text = "星河", "42"
    content = io.BytesIO()
    document.save(content)
    parsed = parse_bytes(content.getvalue(), "application/octet-stream", "client.docx")
    assert "# 客户案例" in parsed.text and "| 星河 | 42 |" in parsed.text


def test_html_parser_removes_scripts_and_navigation():
    parsed = parse_bytes(b"<html><nav>menu</nav><main><h1>Acme</h1><p>Revenue 42</p></main><script>evil()</script></html>",
                         "text/html")
    assert "Revenue 42" in parsed.text
    assert "evil" not in parsed.text and "menu" not in parsed.text


@pytest.mark.parametrize("data,mime,name,error", [
    (b"", "text/plain", "a.txt", "empty_document"),
    (b"\xff\xfe\xfa", "text/plain", "a.txt", "utf8"),
    (b"legacy", "application/msword", "a.doc", "unsupported"),
])
def test_invalid_documents(data, mime, name, error):
    with pytest.raises(ParseError, match=error):
        parse_bytes(data, mime, name)


def test_rrf_and_duplicate_content():
    a = {"id": "a", "text": "one"}
    b = {"id": "b", "text": "two"}
    c = {"id": "c", "text": "one"}
    ranked = fuse({"dense": [a, b], "bm25": [b, c]}, 10)
    assert ranked[0]["id"] == "b"
    assert ranked[0]["fusion_score"] == pytest.approx(1 / 62 + 1 / 61)
    assert len(ranked) == 2


@pytest.mark.asyncio
async def test_ingest_retrieve_citation_and_dedup(library):
    doc = await add(library, "# Customer\n\nAcme lowered acquisition cost to 42 dollars.")
    duplicate = await add(library, "# Customer\n\nAcme lowered acquisition cost to 42 dollars.")
    assert duplicate["duplicate"] and duplicate["document_id"] == doc["document_id"]
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme acquisition"))
    assert result.citations and result.lane_results["bm25"]
    ref = result.citations[0]
    original = await store.citation(library["account"], ref.id)
    assert original["source_text"][ref.start:ref.end] == ref.quote
    output, refs = await validate_output("Cost is 42 [K1]. Invalid [K999].", result, library["account"])
    assert "/api/v1/knowledge/citations/" in output and "[K999]" not in output
    assert refs


@pytest.mark.asyncio
async def test_project_shared_and_account_isolation(library):
    doc = await add(library, "Acme confidential customer evidence.")
    a, p2 = library["account"], library["project2"]
    result = await retrieve(RetrievalRequest(a, p2, "Acme"))
    assert not result.citations
    with pytest.raises(ValueError, match="project_not_found"):
        await retrieve(RetrievalRequest(library["other"], library["project"], "Acme"))
    await store.execute("UPDATE rag_documents SET scope='account' WHERE id=?", (doc["document_id"],))
    result = await retrieve(RetrievalRequest(a, p2, "Acme"))
    assert result.citations  # SQL lanes see sharing immediately, even before payload refresh.
    assert not await store.citation(library["other"], result.citations[0].id)
    await store.execute("UPDATE rag_documents SET scope='project' WHERE id=?", (doc["document_id"],))
    assert not await store.citation(a, result.citations[0].id)


@pytest.mark.asyncio
async def test_external_usage_filter_and_revocation(library):
    doc = await add(library, "Acme private revenue.")
    request = RetrievalRequest(library["account"], library["project"], "Acme", purpose="content")
    assert not (await retrieve(request)).citations
    await store.execute("UPDATE rag_documents SET external_use=1 WHERE id=?", (doc["document_id"],))
    result = await retrieve(request)
    assert result.citations
    await store.execute("UPDATE rag_documents SET external_use=0 WHERE id=?", (doc["document_id"],))
    assert not await store.citation(library["account"], result.citations[0].id)


@pytest.mark.asyncio
async def test_deleted_document_cannot_reappear_from_stale_vectors(library):
    doc = await add(library, "Acme secret source.")
    request = RetrievalRequest(library["account"], library["project"], "Acme")
    result = await retrieve(request)
    await store.execute("UPDATE rag_documents SET deleted=1 WHERE id=?", (doc["document_id"],))
    assert not (await retrieve(request)).citations
    assert not await store.citation(library["account"], result.citations[0].id)
    await ingestion.cleanup(library["account"])


@pytest.mark.asyncio
async def test_failed_new_version_keeps_old_source(library, monkeypatch):
    doc = await add(library, "Acme old evidence.")
    updated = await ingestion.create_document(library["account"], project_id=library["project"],
        title="Updated", data=b"Acme new evidence.", document_id=doc["document_id"])
    monkeypatch.setattr(providers, "embed", AsyncMock(side_effect=RuntimeError("embedding_unavailable")))
    with pytest.raises(RuntimeError):
        await ingestion.ingest(library["account"], doc["document_id"], updated["version_id"])
    row = await store.document(library["account"], doc["document_id"])
    assert row["active_version_id"] == doc["version_id"]


@pytest.mark.asyncio
async def test_rerank_failure_marks_fusion_fallback(library, monkeypatch):
    await add(library, "Acme published case study.")
    monkeypatch.setattr(providers, "rerank", AsyncMock(side_effect=RuntimeError("rerank_unavailable")))
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))
    assert result.status == "degraded" and result.citations
    assert "rerank_unavailable_using_fusion" in result.warnings


@pytest.mark.asyncio
async def test_message_citations_persist_and_check_access(library):
    from opencmo.web import chat_sessions
    await add(library, "Acme research evidence.")
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))
    text, _ = await validate_output("Acme [K1]", result, library["account"])
    session_id = await chat_sessions.create_session(library["project"], library["account"])
    await chat_sessions.update_session(session_id, [{"role": "assistant", "content": text}], library["account"])
    await store.put_message_evidence(session_id, text, result.retrieval_id)
    messages = await chat_sessions.get_session_messages(session_id, library["account"])
    assert messages[0]["citations"]
    assert not await store.message_evidence(library["other"], session_id, text)


def test_handoff_clears_all_private_history():
    from agents import HandoffInputData
    data = HandoffInputData(input_history=({"role": "assistant", "content": "private"},),
                            pre_handoff_items=(), new_items=())
    filtered = safe_handoff_filter("write a post", "public evidence")(data)
    assert "private" not in str(filtered)
    assert content_request("帮我写一个小红书笔记")


@pytest.mark.asyncio
async def test_private_url_and_redirect_targets_rejected(monkeypatch):
    import socket

    from opencmo.rag.fetching import public_target
    loop = __import__("asyncio").get_running_loop()
    monkeypatch.setattr(loop, "getaddrinfo", AsyncMock(return_value=[(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1", 80))]))
    with pytest.raises(ValueError, match="private_url"):
        await public_target("https://rebind.example")


@pytest.mark.asyncio
async def test_knowledge_api_auth_upload_versions_settings(library):
    import httpx

    from opencmo.web.app import app
    token, _ = await storage.create_session(library["user"])
    foreign_token, _ = await storage.create_session(library["other_user"])
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver",
                                cookies={storage.SESSION_COOKIE_NAME: token}) as client:
        config = await client.get("/api/v1/knowledge/settings")
        assert config.status_code == 200
        assert "test-embedding" not in config.text
        created = await client.post("/api/v1/knowledge/documents",
            data={"project_id": str(library["project"]), "title": "Customer document"},
            files={"file": ("customer.md", b"# Customer\n\nAcme customer research.", "text/markdown")})
        assert created.status_code == 202, created.text
        data = created.json()
        await ingestion.ingest(library["account"], data["document_id"], data["version_id"])
        detail = await client.get(f"/api/v1/knowledge/documents/{data['document_id']}")
        assert detail.json()["versions"]
        version = await client.get(f"/api/v1/knowledge/documents/{data['document_id']}/versions/{data['version_id']}")
        assert "Acme" in version.json()["text"]
        assert version.json()["chunks"]
        task = await client.get(f"/api/v1/tasks/{data['task_id']}")
        assert task.status_code == 200
        client.cookies.set(storage.SESSION_COOKIE_NAME, foreign_token)
        assert (await client.get(f"/api/v1/knowledge/documents/{data['document_id']}")).status_code == 404
        assert (await client.get(f"/api/v1/tasks/{data['task_id']}")).status_code == 404
        assert (await client.get(f"/api/v1/tasks/{data['task_id']}/events")).status_code == 404


@pytest.mark.asyncio
async def test_body_limit_rejects_before_parsing(library):
    import httpx

    from opencmo.web.app import app
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as client:
        response = await client.post("/api/v1/knowledge/documents", content=b"{}",
                                     headers={"Content-Length": str(30 * 1024 * 1024), "Content-Type": "application/json"})
        assert response.status_code == 413


@pytest.mark.asyncio
async def test_shadow_rebuild_switches_only_after_success(library, monkeypatch):
    doc = await add(library, "Acme model migration evidence.")
    old = await store.active_profile(library["account"])
    settings = library["settings"].model_copy(update={"embedding_model": "replacement-model"})
    await storage.set_account_setting(library["account"], "RAG_CONFIG", settings.model_dump_json())
    original = providers.embed
    calls = 0
    async def fail_after_probe(texts, settings, **kwargs):
        nonlocal calls
        calls += 1
        if calls > 1:
            raise RuntimeError("embedding_unavailable")
        return await original(texts, settings, **kwargs)
    monkeypatch.setattr(providers, "embed", fail_after_probe)
    with pytest.raises(RuntimeError):
        await ingestion.rebuild(library["account"])
    assert (await store.active_profile(library["account"]))["id"] == old["id"]
    assert (await store.document(library["account"], doc["document_id"]))["active_version_id"] == doc["version_id"]
    monkeypatch.setattr(providers, "embed", original)
    await ingestion.rebuild(library["account"])
    assert (await store.active_profile(library["account"]))["id"] != old["id"]
    assert (await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))).citations


@pytest.mark.asyncio
async def test_invalid_quotation_is_removed(library):
    await add(library, "Acme grew to 42 customers.")
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))
    output, _ = await validate_output('Claim: "Acme grew to 999 customers." [K1]', result, library["account"])
    assert "999" not in output
    assert "unsupported_quotation_removed" in result.warnings


@pytest.mark.asyncio
async def test_low_rerank_scores_produce_no_evidence(library, monkeypatch):
    await add(library, "Acme annual customer report.")
    monkeypatch.setattr(providers, "rerank", AsyncMock(return_value=[(0, 0.01)]))
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))
    assert not result.citations and not result.context


@pytest.mark.asyncio
async def test_no_evidence_chat_does_not_call_agent(library, monkeypatch):
    import agents
    import httpx

    from opencmo.web import chat_sessions
    from opencmo.web.app import app
    token, _ = await storage.create_session(library["user"])
    session = await chat_sessions.create_session(library["project"], library["account"])
    runner = __import__("unittest.mock", fromlist=["Mock"]).Mock(side_effect=AssertionError("must not generate"))
    monkeypatch.setattr(agents.Runner, "run_streamed", runner)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver",
                                cookies={storage.SESSION_COOKIE_NAME: token}) as client:
        response = await client.post("/api/v1/chat", json={"session_id": session, "project_id": library["project"],
            "message": "客户资料有什么结论？", "rag_mode": "only", "locale": "zh"})
        assert response.status_code == 200
        assert '"type": "done"' in response.text
        assert not runner.called


@pytest.mark.asyncio
async def test_native_qdrant_server(library, monkeypatch):
    import os
    import uuid

    from qdrant_client import AsyncQdrantClient
    url = os.environ.get("OPENCMO_RAG_TEST_QDRANT")
    if not url:
        pytest.skip("Set OPENCMO_RAG_TEST_QDRANT for actual server verification.")
    prefix = "test_" + uuid.uuid4().hex[:12]
    monkeypatch.setenv("OPENCMO_QDRANT_PREFIX", prefix)
    monkeypatch.setattr(providers.vectors, "client", lambda: AsyncQdrantClient(url=url))
    collection = None
    try:
        await add(library, "Acme production search engine verification.")
        collection = (await store.active_profile(library["account"]))["collection_name"]
        result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))
        assert result.lane_results["dense"], result.warnings
        assert result.citations
        assert not (await retrieve(RetrievalRequest(library["account"], library["project2"], "Acme"))).citations
    finally:
        if collection:
            client = AsyncQdrantClient(url=url)
            await client.delete_collection(collection)
            await client.close()


@pytest.mark.asyncio
async def test_account_index_lease_serializes_tasks(library, monkeypatch):
    import asyncio
    from types import SimpleNamespace
    active, peak = 0, 0
    async def operation(_ctx):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(.03)
        active -= 1
    monkeypatch.setattr(ingestion, "_execute_knowledge_task", operation)
    contexts = [SimpleNamespace(task={"task_id": f"lease-{i}", "payload": {"account_id": library["account"]}}) for i in range(3)]
    await asyncio.gather(*(ingestion.run_knowledge_executor(ctx) for ctx in contexts))
    assert peak == 1
    assert not await store.one("SELECT * FROM rag_index_leases WHERE account_id=?", (library["account"],))


@pytest.mark.asyncio
async def test_reimport_after_delete_is_still_idempotent(library):
    first = await add(library, "Acme repeated source.")
    await store.execute("UPDATE rag_documents SET deleted=1 WHERE id=?", (first["document_id"],))
    second = await add(library, "Acme repeated source.")
    third = await add(library, "Acme repeated source.")
    assert second["document_id"] != first["document_id"]
    assert third["duplicate"] and third["document_id"] == second["document_id"]


@pytest.mark.asyncio
async def test_project_deletion_immediately_hides_private_documents(library):
    source = await add(library, "Acme deleted project source.")
    await storage.delete_project(library['project'], account_id=library['account'])
    assert await store.document(library["account"], source["document_id"]) is None


@pytest.mark.asyncio
async def test_reindex_creates_shadow_version_and_preserves_source_date(library):
    import httpx

    from opencmo.web.app import app
    item = await add(library, "Acme historical observation.", created_at="2026-01-31 00:00:00")
    token, _ = await storage.create_session(library["user"])
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver",
                                cookies={storage.SESSION_COOKIE_NAME: token}) as client:
        response = await client.post(f"/api/v1/knowledge/documents/{item['document_id']}/reindex")
    assert response.status_code == 200, response.text
    new_version = response.json()["version_id"]
    assert new_version != item["version_id"]
    assert (await store.document(library["account"], item["document_id"]))["active_version_id"] == item["version_id"]
    version = await store.one("SELECT * FROM rag_versions WHERE id=?", (new_version,))
    assert version["created_at"] == "2026-01-31 00:00:00"
    await ingestion.ingest(library["account"], item["document_id"], new_version)
    assert (await store.document(library["account"], item["document_id"]))["active_version_id"] == new_version


@pytest.mark.asyncio
async def test_rebuild_restores_missing_vector_collection(library):
    await add(library, "Acme restore original source.")
    profile = await store.active_profile(library["account"])
    client = providers.vectors.client()
    await client.delete_collection(profile["collection_name"])
    await ingestion.rebuild(library["account"])
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))
    assert result.lane_results["dense"]


@pytest.mark.asyncio
async def test_temporal_filter_applies_to_dense_and_lexical(library):
    await add(library, "Acme past result 42.", title="Acme old", created_at="2026-01-31 00:00:00")
    await add(library, "Acme future result 84.", title="Acme new", created_at="2026-03-31 00:00:00")
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme",
                                             date_from="2026-01-01", date_to="2026-01-31"))
    assert result.citations
    assert all(c.created_at.startswith("2026-01") for c in result.citations)
    assert all(h.title == "Acme old" for h in result.hits)


def test_evaluation_dataset_and_metrics():
    import json
    from pathlib import Path

    from opencmo.rag.evaluate import ranking_metrics, validate_dataset
    dataset = json.loads(Path("tests/fixtures/rag_evaluation.json").read_text(encoding="utf-8"))
    validate_dataset(dataset)
    assert len(dataset["questions"]) == 120
    assert ranking_metrics(["a", "b"], ["b"])["recall_at_20"] == 1
    assert ranking_metrics(["b"], ["b"])["ndcg_at_10"] == 1


def test_pdf_pages_and_ocr_detection():
    from pathlib import Path
    parsed = parse_bytes(Path("tests/fixtures/rag-text.pdf").read_bytes(), "application/pdf")
    assert "42" in parsed.text and "64" in parsed.text
    assert {block.page for block in parsed.blocks} == {1, 2}
    with pytest.raises(ParseError, match="ocr_required"):
        parse_bytes(Path("tests/fixtures/rag-image-only.pdf").read_bytes(), "application/pdf")


@pytest.mark.asyncio
async def test_revoked_evidence_is_not_reused_as_chat_history(library):
    from opencmo.rag.integration import filter_chat_history
    from opencmo.web import chat_sessions
    doc = await add(library, "Acme historical private evidence.")
    result = await retrieve(RetrievalRequest(library["account"], library["project"], "Acme"))
    text, _ = await validate_output("Private evidence [K1].", result, library["account"])
    session_id = await chat_sessions.create_session(library["project"], library["account"])
    await store.put_message_evidence(session_id, text, result.retrieval_id)
    history = [{"role": "assistant", "content": text}]
    assert await filter_chat_history(library["account"], session_id, history, "internal", library["project"], library["project"])
    await store.execute("UPDATE rag_documents SET deleted=1 WHERE id=?", (doc["document_id"],))
    assert not await filter_chat_history(library["account"], session_id, history, "internal", library["project"], library["project"])
    assert not await filter_chat_history(library["account"], session_id, history, "content", library["project"], library["project"])


@pytest.mark.asyncio
async def test_orphan_reconciliation_preserves_foreign_points(library):
    import uuid
    await add(library, "Acme valid evidence.")
    profile = await store.active_profile(library["account"])
    orphan, foreign = str(uuid.uuid4()), str(uuid.uuid4())
    await providers.vectors.upsert(profile["collection_name"], [
        {"id": orphan, "vector": [1.0] * 8, "payload": {"account_id": str(library["account"])}},
        {"id": foreign, "vector": [1.0] * 8, "payload": {"account_id": str(library["other"])}},
    ])
    await ingestion.cleanup(library["account"])
    client = providers.vectors.client()
    remaining = await client.retrieve(profile["collection_name"], [orphan, foreign])
    assert [str(point.id) for point in remaining] == [foreign]


@pytest.mark.asyncio
async def test_provider_usage_metadata_and_invalid_vectors(library, monkeypatch):
    real_embed = library['embed_impl']
    monkeypatch.setattr(providers, "api_post", AsyncMock(return_value={
        "data": [{"index": 0, "embedding": [0.2, 0.4]}], "usage": {"prompt_tokens": 7}, "meta": None}))
    result = await real_embed(["test"], library["settings"], account_id=library["account"])
    assert result == [[.2, .4]]
    event = await store.one("SELECT metadata FROM usage_events WHERE account_id=? AND event_type='rag_embedding' ORDER BY id DESC LIMIT 1",
                            (library["account"],))
    assert event is not None
