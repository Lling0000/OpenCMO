"""Bounded document parsing with source offsets; never silently truncates."""
from __future__ import annotations

import io
import re
import zipfile

from opencmo.rag.types import ParsedDocument, TextBlock

MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 500

class ParseError(ValueError):
    pass

def structured_text(text: str) -> ParsedDocument:
    from markdown_it import MarkdownIt
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    if not text.strip():
        raise ParseError("empty_document")
    lines = text.splitlines(keepends=True)
    offsets = [0]
    for line in lines:
        offsets.append(offsets[-1] + len(line))
    heading_stack: list[tuple[int, str]] = []
    blocks = []
    for token in MarkdownIt("commonmark").enable("table").parse(text):
        if token.map is None or token.level != 0 or token.type.endswith("_close"):
            continue
        start_line, end_line = token.map
        start, end = offsets[start_line], offsets[min(end_line, len(lines))]
        if token.type == "heading_open":
            level = int(token.tag[1:])
            title = re.sub(r"^#+\s*", "", text[start:end]).strip().strip("#").strip()
            heading_stack = [(n, h) for n, h in heading_stack if n < level] + [(level, title)]
        kind = "table" if token.type == "table_open" else ("code" if token.type in {"fence", "code_block"} else "paragraph")
        prefix = "".join(lines[start_line:min(start_line + 2, end_line)]) if kind == "table" else ""
        blocks.append(TextBlock(start, end, " / ".join(h for _, h in heading_stack), kind=kind, prefix=prefix))
    if not blocks:
        blocks = [TextBlock(0, len(text))]
    result = []
    for block in sorted(blocks, key=lambda b: (b.start, -b.end)):
        if result and block.start < result[-1].end:
            continue
        result.append(block)
    return ParsedDocument(text, result)

def _assembled(parts: list[tuple[str, int | None]]) -> ParsedDocument:
    text = ""
    blocks = []
    for content, page in parts:
        parsed = structured_text(content)
        shift = len(text)
        blocks.extend(TextBlock(b.start + shift, b.end + shift, b.heading, page, b.kind, b.prefix) for b in parsed.blocks)
        text += parsed.text + "\n\n"
    if not text.strip():
        raise ParseError("ocr_required_or_empty_document")
    return ParsedDocument(text, blocks)

def _table(rows: list[list]) -> str:
    cleaned = [[str(c or "").replace("\n", " ").replace("|", "\\|") for c in row] for row in rows]
    if not cleaned:
        return ""
    width = max(map(len, cleaned))
    cleaned = [r + [""] * (width - len(r)) for r in cleaned]
    def render(row):
        return "| " + " | ".join(row) + " |"
    return "\n".join([render(cleaned[0]), render(["---"] * width)] + [render(r) for r in cleaned[1:]])

def parse_bytes(data: bytes, mime: str, filename: str = "") -> ParsedDocument:
    if len(data) > MAX_BYTES:
        raise ParseError("file_too_large")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if mime == "application/pdf" or suffix == "pdf":
        import pdfplumber
        parts = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PAGES:
                raise ParseError("pdf_page_limit")
            for number, page in enumerate(pdf.pages, 1):
                tables = page.find_tables()
                boxes = [t.bbox for t in tables]
                outside = page.filter(lambda obj: not any(
                    box[0] <= obj.get("x0", -1) <= box[2] and box[1] <= obj.get("top", -1) <= box[3]
                    for box in boxes
                ))
                body = outside.extract_text(layout=False) or ""
                table_text = "\n\n".join(_table(t.extract()) for t in tables)
                content = "\n\n".join(x for x in (body, table_text) if x.strip())
                if not content.strip():
                    if page.images:
                        raise ParseError(f"ocr_required:page_{number}")
                    continue
                parts.append((content, number))
        return _assembled(parts)
    if suffix == "docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 10000 or sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ParseError("expanded_document_too_large")
        doc = Document(io.BytesIO(data))
        lines = []
        for element in doc.element.body:
            if element.tag.endswith("}p"):
                para = Paragraph(element, doc)
                style = para.style.name if para.style else ""
                match = re.match(r"Heading (\d+)", style)
                lines.append(("#" * min(6, int(match[1])) + " " if match else "") + para.text)
            elif element.tag.endswith("}tbl"):
                table = Table(element, doc)
                lines.append(_table([[c.text for c in row.cells] for row in table.rows]))
        return structured_text("\n\n".join(lines))
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParseError("text_must_be_utf8") from exc
    if mime.startswith("text/html") or suffix in {"html", "htm"}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "noscript", "form"]):
            tag.decompose()
        main = soup.find("main") or soup.find("article") or soup.body or soup
        lines = []
        for tag in main.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "table"]):
            if any(p.name in {"li", "pre", "table"} for p in tag.parents if p is not main):
                continue
            if tag.name == "table":
                lines.append(_table([[c.get_text(" ", strip=True) for c in r.find_all(["th", "td"])] for r in tag.find_all("tr")]))
            elif tag.name == "pre":
                lines.append(chr(96) * 3 + "\n" + tag.get_text() + "\n" + chr(96) * 3)
            else:
                prefix = "#" * int(tag.name[1]) + " " if re.fullmatch(r"h[1-6]", tag.name) else ""
                lines.append(prefix + tag.get_text(" ", strip=True))
        text = "\n\n".join(lines) or main.get_text("\n", strip=True)
    elif suffix not in {"md", "markdown", "txt", ""} and mime not in {"text/plain", "text/markdown"}:
        raise ParseError("unsupported_document_type")
    return structured_text(text)

